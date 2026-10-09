"""Generates Project_Activity_1_Report.docx (run with the project venv)."""
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

REPO = "https://github.com/Kaifovichhh/devops-task-manager"
ACCENT = RGBColor(0x0B, 0x5C, 0x6B)

doc = Document()
st = doc.styles["Normal"]
st.font.name = "Calibri"
st.font.size = Pt(11)
for lvl, size in ((1, 15), (2, 12.5)):
    h = doc.styles[f"Heading {lvl}"]
    h.font.name = "Calibri"
    h.font.size = Pt(size)
    h.font.color.rgb = ACCENT


def shade(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = ""
        r = c.paragraphs[0].add_run(h)
        r.bold = True
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        shade(c, "0B5C6B")
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = v
    if widths:
        from docx.shared import Cm
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    doc.add_paragraph()
    return t


def para(text, bold_prefix=None):
    p = doc.add_paragraph()
    if bold_prefix:
        p.add_run(bold_prefix).bold = True
    p.add_run(text)
    return p


def bullets(items):
    for it in items:
        if isinstance(it, tuple):
            p = doc.add_paragraph(style="List Bullet")
            p.add_run(it[0]).bold = True
            p.add_run(it[1])
        else:
            doc.add_paragraph(it, style="List Bullet")


def code(text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.name = "Courier New"
    r.font.size = Pt(9)
    p.paragraph_format.left_indent = Pt(12)


title = doc.add_heading("Project Activity 1 – Goal, Requirements, GitHub, Features and Docker", 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
para("DevOps", "Course: ")
para("Team TaskOps", "Team Name: ")
para("Task Manager API — a containerised REST service for team task management", "Project: ")
para("Student 1, Student 2, Student 3", "Team members: ")

doc.add_heading("Project Goal", 1)
para("Small student and IT teams often track work in chats and spreadsheets, so tasks get lost, "
     "statuses are unclear and nobody knows the overall progress. Task Manager API gives the team "
     "one place to create, assign a priority to, update and close tasks, and to see progress "
     "statistics. Because it is a simple HTTP API in a Docker container, it can be started on any "
     "machine with one command and integrated with bots, web or mobile front-ends.",
     "Business value: ")
para("", "Objectives our team needed to achieve:")
bullets([
    "Implement a working REST API with CRUD operations, filtering and statistics.",
    "Organise collaboration in GitHub: protected main branch, develop branch, one feature branch per member, Pull Requests.",
    "Follow a test-driven approach: every feature is covered by automated pytest tests (17 tests, 97 % coverage).",
    "Containerise the application with a multi-stage Dockerfile and run it with Docker Compose with persistent data.",
    "Automate lint → tests → Docker build → smoke test with a GitHub Actions CI pipeline.",
    "Document the project so that another team can continue it (README, Git workflow, test cases, backlog).",
])

doc.add_heading("Project user stories", 1)
table(["User Story", "Description"], [
    ("US-1 Create task", "As a team member, I want to create a task with a title, description and priority, so that work is recorded in one place."),
    ("US-2 List & filter", "As a team member, I want to see all tasks and filter them by status (todo / in_progress / done), so that I know what to work on next."),
    ("US-3 Update task", "As a team member, I want to change a task’s status, title or priority, so that the board always reflects reality."),
    ("US-4 Delete task", "As a team lead, I want to delete wrong or duplicate tasks, so that the list stays clean."),
    ("US-5 Statistics", "As a team lead, I want to see how many tasks are in each status, so that I can track progress."),
    ("US-6 Health check", "As a DevOps engineer, I want a /health endpoint, so that Docker and CI can check that the service is alive."),
    ("US-7 Persistent data", "As a user, I want my tasks to survive a container restart, so that no data is lost."),
    ("US-8 One-command run", "As a developer, I want to start the whole app with docker compose up, so that onboarding takes minutes."),
    ("US-9 Automated CI", "As a team, we want every push to be linted, tested and built automatically, so that broken code never reaches main."),
], widths=[4, 12.5])

doc.add_heading("Team member roles and their goals", 1)
table(["Team member", "Role and assigned tasks"], [
    ("Student 1", "Team Lead / Backend developer. Branch feature/tasks-api. Designs the REST API, implements endpoints (US-1…US-6) and input validation in app/routes.py, application factory in app/__init__.py; reviews and merges Pull Requests; protects main."),
    ("Student 2", "QA engineer / Database. Branch feature/storage-tests. Implements the SQLite storage layer app/storage.py (US-7), writes the test plan (docs/TEST_CASES.md) and automated pytest tests, controls coverage ≥ 80 %."),
    ("Student 3", "DevOps engineer. Branch feature/docker-ci. Writes the multi-stage Dockerfile, docker-compose.yml with a volume and healthcheck (US-8), GitHub Actions pipeline (US-9), .dockerignore; prepares the Docker part of the presentation."),
], widths=[3, 13.5])

doc.add_heading("Proof of concept (test-driven approach)", 1)
para("Each member first wrote the test that describes the expected behaviour, saw it fail (red), "
     "implemented the code until it passed (green) and then refactored. All tests run automatically "
     "in GitHub Actions on every push and inside the Docker build.")
bullets([
    ("Student 1 (API): ", "API tests with the Flask test client — e.g. TC-02 POST /tasks returns 201 and status “todo”; TC-03/04 invalid input returns 400; TC-08/11 unknown id returns 404; TC-09 PATCH changes status to “done”. Success = all API tests green."),
    ("Student 2 (Storage / QA): ", "Unit tests on a temporary database — TC-13 data survives re-opening the DB file (simulates a container restart); TC-14 unknown fields are ignored on update; TC-15 missing task returns None/False. Success = tests green and coverage ≥ 80 % (actual 97 %)."),
    ("Student 3 (DevOps): ", "Infrastructure checks — TC-D1 docker build fails if a test fails (tests run in the build stage); TC-D2 container becomes “healthy”; TC-D3 tasks remain after docker compose down/up; TC-D4 container runs as non-root user; TC-D5 CI smoke test creates a task through curl. Success = green GitHub Actions run."),
])

doc.add_heading("Using GitHub for Collaboration", 1)
para(REPO, "Link to the GitHub repository: ")
doc.add_heading("Describe/Create the master branch", 2)
para("The Team Lead created the repository and the main (master) branch with the initial "
     "skeleton (README, .gitignore, requirements). main contains only stable, released code and is "
     "protected in Settings → Branches: direct pushes are forbidden, a Pull Request with one approval "
     "and a green “Lint & Test” CI check are required. Release v1.0.0 is marked with a Git tag.")
code("git init -b main\ngit add README.md .gitignore requirements.txt\n"
     "git commit -m \"chore: initial project skeleton\"\n"
     f"git remote add origin {REPO}.git\ngit push -u origin main\n"
     "git checkout -b develop && git push -u origin develop")
doc.add_heading("Describe/Create your branch", 2)
para("We use a simplified Git Flow. Every member creates a feature branch from develop, commits "
     "with Conventional Commit messages (feat:, test:, ci:, docs:), pushes it and opens a Pull Request "
     "into develop. After review and green CI the branch is merged with --no-ff, then develop is merged "
     "into main for a release.")
code("git checkout develop && git pull\ngit checkout -b feature/docker-ci\n"
     "git add Dockerfile docker-compose.yml .github/\n"
     "git commit -m \"ci: multi-stage Dockerfile, compose and GitHub Actions pipeline\"\n"
     "git push -u origin feature/docker-ci   # then open a Pull Request -> develop")
table(["Branch", "Owner", "Content"], [
    ("main", "Team Lead", "Stable releases (tag v1.0.0)"),
    ("develop", "All", "Integration branch"),
    ("feature/tasks-api", "Student 1", "app/routes.py, app/__init__.py, wsgi.py"),
    ("feature/storage-tests", "Student 2", "app/storage.py, tests/"),
    ("feature/docker-ci", "Student 3", "Dockerfile, docker-compose.yml, .github/workflows/ci.yml"),
], widths=[4.5, 3, 9])

doc.add_heading("Describe your team's application", 1)
doc.add_heading("Features chosen from the backlog and personal tasks", 2)
para("For Sprint 1 we chose the minimum set of features that makes the service useful and lets us "
     "practise the whole DevOps cycle: CRUD for tasks, status filter, statistics, health check, "
     "persistent storage, Docker and CI. Authentication, PostgreSQL, CD and monitoring were moved to "
     "the backlog.")
bullets([
    ("Student 1 – code: ", "REST endpoints GET/POST /tasks, GET/PATCH/DELETE /tasks/<id>, GET /stats, GET /health, JSON validation and JSON error responses. Expected result: API answers with correct codes 200/201/204/400/404."),
    ("Student 2 – code + documents: ", "SQLite repository (create, get, list, update, delete, stats), test plan with 15 test cases and their automation. Expected result: data is stored safely and every feature is covered by tests."),
    ("Student 3 – code + documents: ", "Dockerfile, docker-compose.yml, CI workflow, run instructions in README. Expected result: one command starts the app; every push is checked automatically."),
])
doc.add_heading("Specific objectives and how we test them", 2)
table(["Feature", "Objective", "Test method", "Expected result"], [
    ("Create task", "Valid task is saved, invalid rejected", "pytest API test (TC-02…05)", "201 / 400"),
    ("List & filter", "Return tasks, filter by status", "pytest (TC-06, 07)", "Only matching tasks"),
    ("Update / delete", "Change or remove a task", "pytest (TC-09…11)", "200 / 204 / 404"),
    ("Statistics", "Count tasks per status", "pytest (TC-12)", "Correct counters"),
    ("Persistence", "Data survives restart", "pytest (TC-13) + manual compose down/up", "Task still exists"),
    ("Docker", "Image builds and is healthy", "docker build, docker ps (TC-D1, D2)", "Status “healthy”"),
    ("CI", "Automatic checks on push", "GitHub Actions run (TC-D5)", "All jobs green"),
], widths=[3, 4.5, 5, 4])
para("Testing methods: unit testing (storage layer), API/integration testing with the Flask test "
     "client, static analysis with flake8, code coverage with pytest-cov (threshold 80 %), and a smoke "
     "test of the running container with curl in CI. Local result: 17 passed, coverage 97 %.")
doc.add_heading("Why were your inputs stored? Link to the branch", 2)
para("All inputs (source code, tests, Dockerfile, CI configuration and documentation) are stored in "
     "Git so that: every change has an author and a history and can be rolled back; members work in "
     "parallel in separate branches without breaking each other’s code; Pull Requests give code review; "
     "CI runs automatically on every push; and another team can clone the repository and continue the "
     "project. Task data itself is stored in SQLite inside a Docker volume so that it is not lost when "
     "the container is recreated.")
bullets([
    f"{REPO}/tree/feature/tasks-api",
    f"{REPO}/tree/feature/storage-tests",
    f"{REPO}/tree/feature/docker-ci",
])

doc.add_heading("Strategy of Automation Development of your application", 1)
para("Our strategy is “everything as code, everything automated on push”:")
bullets([
    ("Source control: ", "Git + GitHub, Git Flow branches, protected main, Pull Requests."),
    ("Continuous Integration: ", "GitHub Actions workflow .github/workflows/ci.yml — job 1 installs dependencies, runs flake8 and pytest with coverage ≥ 80 %; job 2 builds the Docker image, starts the container and runs a curl smoke test."),
    ("Quality gate: ", "a red pipeline blocks merging into develop/main; tests also run inside the Docker build, so a broken image can never be produced."),
    ("Infrastructure as code: ", "Dockerfile and docker-compose.yml describe the environment, so “works on my machine” problems disappear."),
    ("Next step (backlog): ", "Continuous Delivery — push the image to GitHub Container Registry on tag v* and deploy automatically to a cloud VM."),
])

doc.add_heading("Container tool, Docker preparation and expected results", 1)
para("We used Docker and Docker Compose.", "Container tool: ")
para("", "Docker preparation:")
bullets([
    ("Multi-stage Dockerfile: ", "stage “test” (python:3.12-slim) installs dev dependencies and runs flake8 + pytest; stage “runtime” contains only Flask + Gunicorn and the app code, which keeps the image small."),
    ("Security: ", "the app runs as non-root user appuser; .dockerignore excludes .git, venv, caches and local DB files."),
    ("Healthcheck: ", "HEALTHCHECK calls /health every 30 s, so Docker shows the container as healthy/unhealthy."),
    ("Persistence: ", "SQLite file /data/tasks.db is stored in the named volume task-data."),
    ("Production server: ", "Gunicorn with 2 workers on port 5000."),
])
code("docker compose up --build -d      # build image and start container\n"
     "docker ps                         # STATUS: Up (healthy)\n"
     "curl http://localhost:5000/health # {\"status\":\"ok\",\"database\":\"ok\",...}\n"
     "curl -X POST http://localhost:5000/tasks -H \"Content-Type: application/json\" \\\n"
     "     -d '{\"title\":\"Prepare slides\"}'\n"
     "docker compose down && docker compose up -d   # task is still there")
para("the image task-manager-api:latest builds successfully only if all tests pass; the container "
     "starts in a few seconds, becomes “healthy”, the API answers on http://localhost:5000, and data "
     "survives container restarts.", "Expected results: ")

doc.add_heading("Future enhancements (backlog)", 1)
bullets([
    "User authentication (JWT) and task assignees",
    "PostgreSQL as a second service in docker-compose",
    "Due dates and overdue notifications",
    "Publish the image to GitHub Container Registry and auto-deploy (CD)",
    "Monitoring with Prometheus + Grafana",
    "Web UI and Kubernetes manifests",
])

doc.add_heading("Reflection", 1)
bullets([
    ("Merge conflicts: ", "two members edited README at the same time. Solution: each member owns separate files and README changes go through one PR."),
    ("Tests not found in Docker build: ", ".dockerignore excluded the tests folder, so the test stage had nothing to run. Solution: keep tests in the build context, but copy only app/ into the runtime stage."),
    ("Data lost after restart: ", "SQLite file was inside the container. Solution: move it to /data and mount a named volume."),
    ("Permissions: ", "the non-root user could not write the DB. Solution: create /data and chown it to appuser in the Dockerfile."),
    ("What we learned: ", "Git Flow and Pull Requests, writing tests before code, multi-stage Docker images, healthchecks and volumes, and building a CI pipeline in GitHub Actions."),
])

doc.save("deliverables/Project_Activity_1_Report.docx")
print("saved")
