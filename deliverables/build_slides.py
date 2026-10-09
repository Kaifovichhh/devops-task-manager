"""Generates Project_Activity_1_Presentation.pptx (10 slides)."""
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

REPO = "github.com/Kaifovichhh/devops-task-manager"
NAVY = RGBColor(0x0F, 0x1E, 0x2E)
TEAL = RGBColor(0x0B, 0x8A, 0x8F)
MINT = RGBColor(0x7F, 0xD6, 0xC2)
ORANGE = RGBColor(0xF2, 0x8C, 0x28)
LIGHT = RGBColor(0xEE, 0xF4, 0xF5)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x1E, 0x29, 0x33)
MUTED = RGBColor(0x5B, 0x6B, 0x78)
HEAD, BODY, MONO = "Cambria", "Calibri", "Courier New"

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]


def bg(slide, color):
    f = slide.background.fill
    f.solid()
    f.fore_color.rgb = color


def text(slide, x, y, w, h, s, size=16, color=INK, bold=False, font=BODY,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, italic=False):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    lines = s if isinstance(s, list) else [s]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(6)
        r = p.add_run()
        r.text = line
        r.font.size, r.font.bold, r.font.italic = Pt(size), bold, italic
        r.font.color.rgb, r.font.name = color, font
    return tb


def box(slide, x, y, w, h, fill, shape=MSO_SHAPE.ROUNDED_RECTANGLE, line=None):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line:
        s.line.color.rgb = line
        s.line.width = Pt(1.25)
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = 0.08
    return s


def badge(slide, x, y, d, label, fill=TEAL, size=16):
    c = box(slide, x, y, d, d, fill, MSO_SHAPE.OVAL)
    tf = c.text_frame
    tf.margin_left = tf.margin_right = 0
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = label
    r.font.size, r.font.bold, r.font.color.rgb, r.font.name = Pt(size), True, WHITE, BODY
    return c


def title(slide, s, kicker=None, dark=False):
    if kicker:
        text(slide, 0.6, 0.45, 12, 0.35, kicker.upper(), 12, TEAL if not dark else MINT, True)
    text(slide, 0.6, 0.75, 12.1, 0.8, s, 34, WHITE if dark else NAVY, True, HEAD)


def notes(slide, s):
    slide.notes_slide.notes_text_frame.text = s


# 1 ─ Title
s = prs.slides.add_slide(BLANK)
bg(s, NAVY)
for i, (x, y) in enumerate([(9.2, 1.2), (10.6, 1.2), (9.2, 2.6), (10.6, 2.6), (11.95, 2.6)]):
    box(s, x, y, 1.2, 1.2, TEAL if i % 2 == 0 else MINT)
text(s, 0.8, 1.3, 8, 0.4, "DEVOPS · PROJECT ACTIVITY 1", 14, MINT, True)
text(s, 0.8, 1.9, 8.2, 1.8, "Task Manager API", 54, WHITE, True, HEAD)
text(s, 0.8, 3.15, 8, 1, "A containerised REST service for team task management — "
     "built with Git Flow, TDD, Docker and GitHub Actions", 20, LIGHT)
text(s, 0.8, 5.3, 8, 0.4, "Team TaskOps", 18, ORANGE, True)
text(s, 0.8, 5.8, 9, 0.4, "Rakhat · Danas · Nurlan", 16, LIGHT)
text(s, 0.8, 6.3, 9, 0.4, REPO, 14, MINT, font=MONO)
notes(s, "Introduce the team and the project in one sentence.")

# 2 ─ Goal & business value
s = prs.slides.add_slide(BLANK)
bg(s, WHITE)
title(s, "One place for team tasks, one command to run it", "Project goal & business value")
text(s, 0.6, 1.9, 5.8, 4.5, [
    "Problem: tasks live in chats and spreadsheets — they get lost and progress is invisible.",
    "Solution: a simple HTTP API to create, prioritise, update, filter and close tasks, plus live statistics.",
    "Value: starts anywhere with Docker in seconds and can be connected to a bot, web or mobile UI.",
    "DevOps objective: practise the full cycle — Git, TDD, Docker, CI.",
], 17, INK)
stats = [("9", "user stories"), ("17", "automated tests"), ("97%", "code coverage"), ("1", "command to run")]
for i, (n, lbl) in enumerate(stats):
    x, y = 7.0 + (i % 2) * 3.0, 1.9 + (i // 2) * 2.4
    box(s, x, y, 2.7, 2.1, LIGHT)
    text(s, x, y + 0.3, 2.7, 1, n, 54, TEAL, True, HEAD, PP_ALIGN.CENTER)
    text(s, x, y + 1.4, 2.7, 0.4, lbl, 15, MUTED, align=PP_ALIGN.CENTER)

# 3 ─ User stories
s = prs.slides.add_slide(BLANK)
bg(s, LIGHT)
title(s, "Nine user stories define Sprint 1", "Requirements")
stories = [
    ("US-1", "Create task", "with title, description, priority"),
    ("US-2", "List & filter", "by status todo / in_progress / done"),
    ("US-3", "Update task", "change status, title, priority"),
    ("US-4", "Delete task", "keep the list clean"),
    ("US-5", "Statistics", "tasks per status for the lead"),
    ("US-6", "Health check", "/health for Docker and CI"),
    ("US-7", "Persistent data", "survives container restart"),
    ("US-8", "One-command run", "docker compose up"),
    ("US-9", "Automated CI", "lint, test, build on every push"),
]
for i, (code_, name, desc) in enumerate(stories):
    x, y = 0.6 + (i % 3) * 4.1, 1.85 + (i // 3) * 1.75
    box(s, x, y, 3.85, 1.5, WHITE)
    badge(s, x + 0.25, y + 0.35, 0.8, code_, TEAL if i < 5 else ORANGE, 12)
    text(s, x + 1.25, y + 0.3, 2.5, 0.4, name, 17, NAVY, True)
    text(s, x + 1.25, y + 0.75, 2.5, 0.6, desc, 13, MUTED)
text(s, 0.6, 7.0, 12, 0.3, "Teal = application features · Orange = DevOps / infrastructure features",
     11, MUTED)

# 4 ─ Team
s = prs.slides.add_slide(BLANK)
bg(s, WHITE)
title(s, "Three roles, three branches, clear ownership", "Team members & roles")
team = [
    ("R", "Rakhat", "Team Lead · Backend", "feature/tasks-api",
     ["REST endpoints & validation", "App factory, error handling", "Code review, merges to main"]),
    ("D", "Danas", "QA · Database", "feature/storage-tests",
     ["SQLite storage layer", "Test plan: 15 test cases", "pytest + coverage ≥ 80 %"]),
    ("N", "Nurlan", "DevOps engineer", "feature/docker-ci",
     ["Multi-stage Dockerfile", "docker-compose + volume", "GitHub Actions pipeline"]),
]
for i, (ini, name, role, br, tasks) in enumerate(team):
    x = 0.6 + i * 4.1
    box(s, x, 1.9, 3.85, 4.9, LIGHT)
    badge(s, x + 1.43, 2.2, 1.0, ini, [TEAL, NAVY, ORANGE][i], 22)
    text(s, x + 0.3, 3.35, 3.25, 0.4, name, 20, NAVY, True, HEAD, PP_ALIGN.CENTER)
    text(s, x + 0.3, 3.8, 3.25, 0.4, role, 15, TEAL, True, align=PP_ALIGN.CENTER)
    text(s, x + 0.3, 4.25, 3.25, 0.4, br, 12, MUTED, font=MONO, align=PP_ALIGN.CENTER)
    text(s, x + 0.4, 4.9, 3.1, 1.8, ["• " + t for t in tasks], 14, INK)

# 5 ─ Application & features
s = prs.slides.add_slide(BLANK)
bg(s, WHITE)
title(s, "A small REST API on Flask + SQLite", "Application & features")
rows = [("GET", "/health", "service & DB status"), ("GET", "/tasks?status=", "list / filter"),
        ("POST", "/tasks", "create task"), ("GET", "/tasks/<id>", "get one task"),
        ("PATCH", "/tasks/<id>", "update / complete"), ("DELETE", "/tasks/<id>", "delete"),
        ("GET", "/stats", "count per status")]
colors = {"GET": TEAL, "POST": ORANGE, "PATCH": NAVY, "DELETE": RGBColor(0xC0, 0x39, 0x2B)}
for i, (m, path, d) in enumerate(rows):
    y = 1.9 + i * 0.68
    b = box(s, 0.6, y, 1.2, 0.5, colors[m])
    tf = b.text_frame
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    r = tf.paragraphs[0].add_run()
    r.text, r.font.size, r.font.bold, r.font.color.rgb = m, Pt(12), True, WHITE
    text(s, 2.0, y + 0.08, 2.6, 0.4, path, 14, INK, True, MONO)
    text(s, 4.7, y + 0.1, 2.6, 0.4, d, 14, MUTED)
# architecture
for i, (lbl, sub) in enumerate([("Client", "curl / UI / bot"), ("Gunicorn + Flask", "routes.py"),
                                ("Storage", "storage.py"), ("SQLite", "volume /data")]):
    y = 1.9 + i * 1.25
    box(s, 8.4, y, 4.3, 0.95, LIGHT if i % 2 else NAVY)
    text(s, 8.6, y + 0.12, 3.9, 0.4, lbl, 16, WHITE if i % 2 == 0 else NAVY, True)
    text(s, 8.6, y + 0.5, 3.9, 0.4, sub, 12, MINT if i % 2 == 0 else MUTED)
    if i < 3:
        a = s.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Inches(10.4), Inches(y + 0.97),
                               Inches(0.3), Inches(0.26))
        a.fill.solid(); a.fill.fore_color.rgb = TEAL; a.line.fill.background()
text(s, 0.6, 6.85, 12, 0.4, "Why these features: the minimum that makes the service useful AND "
     "exercises every DevOps step (data, tests, container, CI).", 13, MUTED, italic=True)

# 6 ─ GitHub
s = prs.slides.add_slide(BLANK)
bg(s, LIGHT)
title(s, "Git Flow: features → develop → protected main", "GitHub collaboration")
lanes = [("main", NAVY, 2.2), ("develop", TEAL, 3.4), ("feature/*", ORANGE, 4.6)]
for name, col, y in lanes:
    text(s, 0.6, y - 0.18, 1.6, 0.4, name, 14, col, True, MONO)
    ln = s.shapes.add_connector(1, Inches(2.3), Inches(y), Inches(9.0), Inches(y))
    ln.line.color.rgb, ln.line.width = col, Pt(3)
pts = [(2.5, 2.2, NAVY), (8.7, 2.2, NAVY), (3.2, 3.4, TEAL), (4.6, 3.4, TEAL), (6.0, 3.4, TEAL),
       (7.4, 3.4, TEAL), (4.0, 4.6, ORANGE), (5.4, 4.6, ORANGE), (6.8, 4.6, ORANGE)]
for x, y, c in pts:
    box(s, x - 0.15, y - 0.15, 0.3, 0.3, c, MSO_SHAPE.OVAL)
text(s, 8.25, 1.65, 1.5, 0.3, "v1.0.0", 12, NAVY, True)
for x, lbl in [(4.0, ["Danas", "storage"]), (5.4, ["Rakhat", "api"]), (6.8, ["Nurlan", "docker-ci"])]:
    text(s, x - 0.7, 4.85, 1.4, 0.5, lbl, 11, MUTED, align=PP_ALIGN.CENTER)
box(s, 9.6, 1.85, 3.15, 3.4, WHITE)
text(s, 9.85, 2.05, 2.7, 3.1, ["Rules on main", "• no direct push", "• PR + 1 review",
                                "• CI must be green", "• releases are tagged"], 14, INK)
box(s, 0.6, 5.55, 12.15, 1.4, NAVY)
text(s, 0.9, 5.7, 11.6, 0.4, "git checkout develop && git checkout -b feature/docker-ci", 14, MINT,
     font=MONO)
text(s, 0.9, 6.1, 11.6, 0.4, "git commit -m \"ci: multi-stage Dockerfile\" && git push -u origin "
     "feature/docker-ci   → Pull Request", 14, WHITE, font=MONO)
text(s, 0.9, 6.5, 11.6, 0.4, REPO, 13, ORANGE, True, MONO)

# 7 ─ Testing
s = prs.slides.add_slide(BLANK)
bg(s, WHITE)
title(s, "Test first: every feature has an automated check", "Proof of concept · TDD")
cases = [("TC-02", "POST /tasks valid", "201, status todo"),
         ("TC-03/04", "missing / blank title", "400"),
         ("TC-07", "filter ?status=done", "only done tasks"),
         ("TC-09", "PATCH status=done", "200, updated"),
         ("TC-11", "DELETE twice", "204, then 404"),
         ("TC-13", "reopen DB file", "data persists"),
         ("TC-D2", "docker ps", "healthy"),
         ("TC-D5", "CI smoke test", "all jobs green")]
tbl = s.shapes.add_table(len(cases) + 1, 3, Inches(0.6), Inches(1.9), Inches(7.6),
                         Inches(4.8)).table
for j, (h, w) in enumerate([("ID", 1.4), ("Scenario", 3.4), ("Expected", 2.8)]):
    tbl.columns[j].width = Inches(w)
    tbl.cell(0, j).text = h
for i, row in enumerate(cases, 1):
    for j, v in enumerate(row):
        tbl.cell(i, j).text = v
for i in range(len(cases) + 1):
    for j in range(3):
        c = tbl.cell(i, j)
        c.fill.solid()
        c.fill.fore_color.rgb = NAVY if i == 0 else (LIGHT if i % 2 else WHITE)
        for p in c.text_frame.paragraphs:
            for r in p.runs:
                r.font.size, r.font.name = Pt(14), BODY
                r.font.color.rgb = WHITE if i == 0 else INK
                r.font.bold = i == 0
for i, (n, lbl, col) in enumerate([("17 / 17", "tests passed", TEAL),
                                    ("97 %", "coverage (gate ≥ 80 %)", ORANGE),
                                    ("0", "flake8 warnings", NAVY)]):
    y = 1.9 + i * 1.65
    box(s, 8.7, y, 4.05, 1.4, LIGHT)
    text(s, 8.95, y + 0.15, 3.6, 0.7, n, 34, col, True, HEAD)
    text(s, 8.95, y + 0.9, 3.6, 0.4, lbl, 14, MUTED)

# 8 ─ Docker
s = prs.slides.add_slide(BLANK)
bg(s, LIGHT)
title(s, "Multi-stage image: tests gate the build", "Docker preparation")
stages = [("Stage 1 · test", "python:3.12-slim", ["install dev deps", "flake8 + pytest", "build fails if a test fails"], TEAL),
          ("Stage 2 · runtime", "python:3.12-slim", ["Flask + Gunicorn only", "non-root appuser", "HEALTHCHECK /health"], NAVY)]
for i, (h, sub, items, col) in enumerate(stages):
    x = 0.6 + i * 4.3
    box(s, x, 1.9, 3.8, 3.3, col)
    text(s, x + 0.3, 2.1, 3.3, 0.4, h, 18, WHITE, True)
    text(s, x + 0.3, 2.55, 3.3, 0.4, sub, 12, MINT, font=MONO)
    text(s, x + 0.3, 3.1, 3.3, 2, ["• " + t for t in items], 15, WHITE)
a = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(4.5), Inches(3.35), Inches(0.35), Inches(0.4))
a.fill.solid(); a.fill.fore_color.rgb = ORANGE; a.line.fill.background()
box(s, 9.2, 1.9, 3.55, 3.3, WHITE)
text(s, 9.45, 2.1, 3.1, 3, ["docker-compose.yml", "• service api :5000", "• volume task-data → /data",
                             "• restart: unless-stopped"], 15, INK)
box(s, 0.6, 5.5, 12.15, 1.5, NAVY)
text(s, 0.9, 5.65, 11.6, 1.3, ["$ docker compose up --build -d",
                               "$ docker ps            →  task-manager-api   Up (healthy)",
                               "$ curl localhost:5000/health  →  {\"status\": \"ok\", \"database\": \"ok\"}"],
     14, MINT, font=MONO)

# 9 ─ Automation strategy
s = prs.slides.add_slide(BLANK)
bg(s, WHITE)
title(s, "Every push runs the same automated pipeline", "Automation strategy · GitHub Actions")
steps = [("1", "Push / PR", "feature/*, develop, main"), ("2", "Lint", "flake8"),
         ("3", "Test", "pytest + coverage ≥ 80 %"), ("4", "Build", "docker build (tests inside)"),
         ("5", "Smoke test", "run container, curl API"), ("6", "Merge", "only if all green")]
for i, (n, h, d) in enumerate(steps):
    x = 0.6 + i * 2.1
    badge(s, x + 0.55, 2.0, 0.8, n, ORANGE if i in (0, 5) else TEAL, 20)
    if i < 5:
        ln = s.shapes.add_connector(1, Inches(x + 1.45), Inches(2.4), Inches(x + 2.55), Inches(2.4))
        ln.line.color.rgb, ln.line.width = MUTED, Pt(2)
    text(s, x, 3.0, 1.9, 0.4, h, 17, NAVY, True, align=PP_ALIGN.CENTER)
    text(s, x, 3.45, 1.9, 0.8, d, 13, MUTED, align=PP_ALIGN.CENTER)
box(s, 0.6, 4.6, 5.9, 2.3, LIGHT)
text(s, 0.9, 4.8, 5.4, 2, ["Everything as code",
                           "Dockerfile, compose and CI YAML live in Git, reviewed like any other code."],
     15, INK)
box(s, 6.85, 4.6, 5.9, 2.3, LIGHT)
text(s, 7.15, 4.8, 5.4, 2, ["Quality gate",
                            "Red pipeline blocks the merge — broken code never reaches main."], 15, INK)

# 10 ─ Backlog & reflection
s = prs.slides.add_slide(BLANK)
bg(s, NAVY)
title(s, "What we learned and what comes next", "Backlog & reflection", dark=True)
box(s, 0.6, 1.9, 5.9, 4.9, RGBColor(0x18, 0x2E, 0x44))
text(s, 0.9, 2.1, 5.4, 0.4, "Future enhancements", 20, ORANGE, True, HEAD)
text(s, 0.9, 2.7, 5.4, 4, ["• JWT auth and task assignees", "• PostgreSQL service in compose",
                           "• Push image to GHCR + auto-deploy (CD)", "• Prometheus + Grafana monitoring",
                           "• Web UI, Kubernetes manifests"], 16, WHITE)
box(s, 6.85, 1.9, 5.9, 4.9, RGBColor(0x18, 0x2E, 0x44))
text(s, 7.15, 2.1, 5.4, 0.4, "Issues → solutions", 20, MINT, True, HEAD)
text(s, 7.15, 2.7, 5.4, 4, ["• Tests missing in build → fixed .dockerignore",
                            "• Data lost on restart → named volume /data",
                            "• Non-root can't write DB → chown /data",
                            "• Parallel edits → file ownership + PRs",
                            "Learned: Git Flow, TDD, multi-stage Docker, CI"], 16, WHITE)
text(s, 0.6, 6.95, 12, 0.35, "Thank you! Questions?   ·   " + REPO, 14, MINT)

prs.save("deliverables/Project_Activity_1_Presentation.pptx")
print("saved", len(prs.slides), "slides")
